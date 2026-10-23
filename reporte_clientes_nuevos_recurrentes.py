import psycopg2
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os
import time

# ============================================================
# CONFIGURACIÓN DE CONEXIÓN
# ============================================================
DB_CONFIG = {
    'host':     '192.168.91.131',
    'port':     5432,
    'dbname':   'cadm_j5004625301',
    'user':     'postgres',
    'password': 'Soloyo1234*-',
    'client_encoding': 'UTF8'
}

# ============================================================
# PERIODO DE ANÁLISIS
# ============================================================
FECHA_INICIO_PERIODO = '2026-09-01'   # << Cambia aquí el inicio del periodo

# ============================================================
# QUERIES AUXILIARES PARA LOS FILTROS
# ============================================================
QUERY_VENDEDORES = """
SELECT DISTINCT TRIM(description) AS vendedor
FROM public.sellers
WHERE code = ANY(ARRAY[
    '25625528', '25411189', '23225823',
    '22974980', '281502778', '24057112',
    '28386523'
])
  AND description IS NOT NULL
  AND TRIM(description) <> ''
ORDER BY 1;
"""

QUERY_PROVINCIAS = """
SELECT DISTINCT 
    TRIM(description) AS provincia
FROM public.provinces
WHERE description IS NOT NULL 
  AND TRIM(description) <> ''
ORDER BY 1;
"""

# ============================================================
# HELPERS DE SELECCIÓN INTERACTIVA
# ============================================================

def seleccionar_multiple(titulo, opciones, etiqueta_todos="TODOS"):
    """
    Muestra una lista numerada y permite al usuario seleccionar
    uno o varios elementos separados por coma, o ENTER para todos.
    Retorna la lista de valores seleccionados, o [] si eligió todos.
    """
    print(f"\n{'='*60}")
    print(f"  {titulo}")
    print(f"{'='*60}")
    print(f"  [0] {etiqueta_todos}")
    for i, op in enumerate(opciones, start=1):
        print(f"  [{i}] {op}")
    print(f"{'='*60}")
    raw = input("Ingresa los números separados por coma (o ENTER para TODOS): ").strip()

    if not raw:
        print(f"  >> Seleccionado: {etiqueta_todos}")
        return []

    seleccionados = []
    for parte in raw.split(','):
        parte = parte.strip()
        if parte == '0':
            print(f"  >> Seleccionado: {etiqueta_todos}")
            return []
        try:
            idx = int(parte)
            if 1 <= idx <= len(opciones):
                seleccionados.append(opciones[idx - 1])
            else:
                print(f"  [!] Índice {idx} fuera de rango, ignorado.")
        except ValueError:
            print(f"  [!] '{parte}' no es un número válido, ignorado.")

    if seleccionados:
        print(f"  >> Seleccionados: {', '.join(seleccionados)}")
    else:
        print(f"  >> Sin selección válida, se usará: {etiqueta_todos}")
    return seleccionados


def seleccionar_tipo_cliente():
    """Selección simple: Nuevo / Recurrente / Todos."""
    opciones = ['NUEVO', 'RECURRENTE', 'TODOS']
    print(f"\n{'='*60}")
    print("  FILTRO: Tipo de Cliente")
    print(f"{'='*60}")
    for i, op in enumerate(opciones, start=1):
        print(f"  [{i}] {op}")
    print(f"{'='*60}")
    raw = input("Elige una opción [1-3] (o ENTER para TODOS): ").strip()
    mapa = {'1': 'NUEVO', '2': 'RECURRENTE', '3': 'TODOS'}
    tipo = mapa.get(raw, 'TODOS')
    print(f"  >> Seleccionado: {tipo}")
    return tipo


# ============================================================
# CONSTRUCCIÓN DINÁMICA DE LA QUERY PRINCIPAL
# ============================================================

def build_query(vendedores_sel, provincias_sel, tipo_cliente, fecha_inicio):
    """
    Construye la query de clientes nuevos/recurrentes con los filtros
    dinámicos aplicados directamente en SQL.
    """

    # --- Filtro de vendedores ---
    if vendedores_sel:
        lista = ", ".join(f"'{v.replace(chr(39), chr(39)*2)}'" for v in vendedores_sel)
        filtro_vendedores_sales_base = f"AND TRIM(sel.description) IN ({lista})"
    else:
        # Sin filtro: incluir los vendedores del equipo (por código)
        filtro_vendedores_sales_base = """AND sel.code IN (
            '25625528', '25411189', '23225823',
            '22974980', '281502778', '24057112',
            '28386523'
        )"""

    # --- Filtro de tipo de cliente ---
    if tipo_cliente == 'NUEVO':
        filtro_tipo = f"WHERE pce.primera_compra_historica >= '{fecha_inicio}'"
    elif tipo_cliente == 'RECURRENTE':
        filtro_tipo = f"WHERE pce.primera_compra_historica < '{fecha_inicio}'"
    else:
        filtro_tipo = ""

    # --- Filtro de provincia ---
    if provincias_sel:
        lista_prov = ", ".join(f"'{p.replace(chr(39), chr(39)*2)}'" for p in provincias_sel)
        filtro_provincia = f"TRIM(p.description) IN ({lista_prov})"
    else:
        filtro_provincia = ""

    # Combinar filtros
    if filtro_tipo and filtro_provincia:
        where_final = filtro_tipo + f" AND {filtro_provincia}"
    elif filtro_tipo:
        where_final = filtro_tipo
    elif filtro_provincia:
        where_final = f"WHERE {filtro_provincia}"
    else:
        where_final = ""

    query = f"""
WITH invoice_units AS (
    SELECT 
        main_correlative,
        SUM(amount) AS total_unidades
    FROM public.sales_operation_details
    GROUP BY main_correlative
),
primera_compra_empresa AS (
    SELECT 
        client_code AS rif,
        MIN(emission_date) AS primera_compra_historica
    FROM public.sales_operation
    WHERE operation_type = 'BILL'
      AND client_name <> 'ANULADO'
      AND user_code = 'ALMAYOR'
    GROUP BY client_code
),
sales_base AS (
    SELECT 
        s.correlative,
        s.document_no AS factura,
        s.emission_date AS fecha,
        s.register_hour AS hora,
        s.client_code AS rif,
        s.client_name,
        s.seller AS seller_code,
        sel.description AS vendedor,
        s.total_net_details AS monto_bs,
        sc.sales_aliquot AS tasa_del_dia,
        ROUND((s.total_net_details / NULLIF(sc.sales_aliquot, 0))::numeric, 2) AS monto_usd,
        COALESCE(iu.total_unidades, 0) AS unidades
    FROM public.sales_operation s
    JOIN public.sales_operation_coins sc 
        ON s.correlative = sc.main_correlative 
       AND sc.coin_code = '02'
    LEFT JOIN invoice_units iu 
        ON s.correlative = iu.main_correlative
    JOIN public.sellers sel 
        ON s.seller = sel.code
    WHERE s.user_code = 'ALMAYOR' 
      AND s.operation_type = 'BILL' 
      AND s.client_name <> 'ANULADO'
      AND s.emission_date >= '{fecha_inicio}'
      {filtro_vendedores_sales_base}
),
ranked_sales AS (
    SELECT 
        sb.*,
        ROW_NUMBER() OVER (
            PARTITION BY sb.rif 
            ORDER BY sb.fecha DESC, sb.hora DESC, sb.correlative DESC
        ) AS rn_ultimo
    FROM sales_base sb
),
client_sales_summary AS (
    SELECT 
        rif,
        COUNT(DISTINCT factura)    AS total_facturas,
        SUM(unidades)              AS cantidad_unidades_compradas,
        SUM(monto_usd)             AS monto_total_dolares,
        SUM(monto_bs)              AS monto_total_bs,
        MIN(fecha)                 AS primera_compra_periodo,
        MAX(fecha)                 AS ultima_compra
    FROM sales_base
    GROUP BY rif
),
ultimo_vendedor AS (
    SELECT 
        rif,
        vendedor AS vendedor_que_le_vendio
    FROM ranked_sales
    WHERE rn_ultimo = 1
)
SELECT 
    css.rif                                                              AS "RIF / Código",
    COALESCE(NULLIF(TRIM(c.description), ''), css.rif)                  AS "Cliente",
    COALESCE(NULLIF(TRIM(p.description), ''), 'NO REGISTRADO')          AS "Provincia / Estado",
    uv.vendedor_que_le_vendio                                            AS "Persona que le Vendió",
    CASE 
        WHEN pce.primera_compra_historica >= '{fecha_inicio}' THEN 'NUEVO'
        ELSE 'RECURRENTE'
    END                                                                  AS "Tipo de Cliente",
    pce.primera_compra_historica                                         AS "Primera Compra Empresa",
    css.primera_compra_periodo                                           AS "Primera Compra Periodo",
    css.ultima_compra                                                    AS "Última Compra",
    css.total_facturas                                                   AS "Total Facturas",
    css.cantidad_unidades_compradas                                      AS "Cantidad Total Unidades Compradas",
    css.monto_total_dolares                                              AS "Monto Total en Dólares ($)",
    css.monto_total_bs                                                   AS "Monto Total en Bs.",
    COALESCE(NULLIF(TRIM(c.phone), ''), 'SIN TELEFONO')                 AS "Teléfono"
FROM client_sales_summary css
LEFT JOIN primera_compra_empresa pce 
    ON css.rif = pce.rif
LEFT JOIN public.clients c 
    ON css.rif = c.code
LEFT JOIN public.provinces p 
    ON c.province = p.code
LEFT JOIN ultimo_vendedor uv 
    ON css.rif = uv.rif
{where_final}
ORDER BY css.monto_total_dolares DESC;
"""
    return query


# ============================================================
# EXPORTACIÓN A EXCEL
# ============================================================

def exportar_excel(df, excel_path, tipo_cliente, vendedores_sel, provincias_sel, fecha_inicio):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Estilos
    header_font  = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill  = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    nuevo_fill   = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")   # verde suave
    recur_fill   = PatternFill(start_color="D0E8FF", end_color="D0E8FF", fill_type="solid")   # azul suave
    alt_fill     = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    border_thin  = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    num_formats = {
        "Total Facturas":                    "#,##0",
        "Cantidad Total Unidades Compradas": "#,##0",
        "Monto Total en Dólares ($)":        "$#,##0.00",
        "Monto Total en Bs.":                "#,##0.00",
    }

    ws = wb.create_sheet(title="Clientes Nuevos y Recurrentes")

    # Fila 1: parámetros usados (informativa)
    info_vendedores = ', '.join(vendedores_sel) if vendedores_sel else 'TODOS LOS VENDEDORES DEL EQUIPO'
    info_provincias = ', '.join(provincias_sel) if provincias_sel else 'TODAS LAS PROVINCIAS'
    ws.append([
        f"Periodo desde: {fecha_inicio}  |  Tipo: {tipo_cliente}  "
        f"|  Vendedores: {info_vendedores}  |  Provincias: {info_provincias}"
    ])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(df.columns))
    info_cell = ws.cell(row=1, column=1)
    info_cell.font = Font(name="Calibri", size=10, italic=True, color="444444")
    info_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 20

    # Fila 2: encabezados
    headers = list(df.columns)
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=2, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_thin
    ws.row_dimensions[2].height = 30

    tipo_col_idx = headers.index("Tipo de Cliente") + 1 if "Tipo de Cliente" in headers else None

    # Datos desde fila 3
    for row_idx, row_data in enumerate(df.itertuples(index=False), start=3):
        ws.append(list(row_data))
        tipo_valor = ws.cell(row=row_idx, column=tipo_col_idx).value if tipo_col_idx else None
        is_alt = (row_idx % 2 == 0)

        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = border_thin

            # Color de fila según tipo de cliente
            if tipo_valor == 'NUEVO':
                cell.fill = nuevo_fill
            elif tipo_valor == 'RECURRENTE':
                cell.fill = recur_fill
            elif is_alt:
                cell.fill = alt_fill

            header_name = headers[col_idx - 1]
            if header_name in num_formats:
                cell.number_format = num_formats[header_name]
                cell.alignment = Alignment(horizontal="right")
            elif isinstance(val, (int, float)):
                cell.alignment = Alignment(horizontal="right")
            elif any(k in header_name for k in ("Fecha", "Compra", "Última")):
                cell.alignment = Alignment(horizontal="center")
            elif header_name == "Tipo de Cliente":
                cell.alignment = Alignment(horizontal="center")
            else:
                cell.alignment = Alignment(horizontal="left")

    # Auto-ajuste de columnas
    for col_idx, header_name in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(header_name)
        for r in range(3, min(ws.max_row + 1, 200)):
            v = ws.cell(row=r, column=col_idx).value
            if v is not None:
                max_len = max(max_len, len(str(v)))
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    # Hoja de resumen por tipo
    ws_res = wb.create_sheet(title="Resumen por Tipo")
    ws_res.append(["Tipo de Cliente", "Cantidad de Clientes", "Monto Total USD ($)"])
    for col_idx in range(1, 4):
        cell = ws_res.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border_thin
    ws_res.row_dimensions[1].height = 28

    resumen = df.groupby("Tipo de Cliente").agg(
        Clientes=("RIF / Código", "count"),
        Total_USD=("Monto Total en Dólares ($)", "sum")
    ).reset_index()

    for row_idx, row in enumerate(resumen.itertuples(index=False), start=2):
        ws_res.append([row[0], row[1], row[2]])
        for col_idx in range(1, 4):
            cell = ws_res.cell(row=row_idx, column=col_idx)
            cell.border = border_thin
            if col_idx == 3:
                cell.number_format = "$#,##0.00"
                cell.alignment = Alignment(horizontal="right")
            elif col_idx == 2:
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.alignment = Alignment(horizontal="center")

    for i, w in enumerate([22, 22, 22], start=1):
        ws_res.column_dimensions[get_column_letter(i)].width = w

    wb.save(excel_path)


# ============================================================
# FUNCIÓN PRINCIPAL
# ============================================================

def generar_reporte():
    inicio_tiempo = time.time()

    print("\n" + "="*60)
    print("  REPORTE: CLIENTES NUEVOS Y RECURRENTES AL MAYOR")
    print("="*60)
    print(f"  Periodo desde: {FECHA_INICIO_PERIODO}")

    # Conexión para cargar opciones de filtro
    print("\nConectando a la base de datos para cargar opciones de filtro...")
    conn = psycopg2.connect(**DB_CONFIG)

    vendedores_opciones = pd.read_sql_query(QUERY_VENDEDORES, conn)['vendedor'].tolist()
    provincias_opciones = pd.read_sql_query(QUERY_PROVINCIAS, conn)['provincia'].tolist()

    # Filtros interactivos
    vendedores_sel = seleccionar_multiple(
        "FILTRO: Vendedores (del equipo)",
        vendedores_opciones,
        etiqueta_todos="TODOS LOS VENDEDORES DEL EQUIPO"
    )

    provincias_sel = seleccionar_multiple(
        "FILTRO: Provincia / Estado",
        provincias_opciones,
        etiqueta_todos="TODAS LAS PROVINCIAS"
    )

    tipo_cliente = seleccionar_tipo_cliente()

    # Ejecutar consulta principal
    print("\nEjecutando consulta principal con los filtros seleccionados...")
    query = build_query(vendedores_sel, provincias_sel, tipo_cliente, FECHA_INICIO_PERIODO)

    df = pd.read_sql_query(query, conn)
    conn.close()
    print(f"Registros obtenidos: {len(df):,} clientes")

    if df.empty:
        print("\n[!] No se encontraron registros con los filtros seleccionados.")
        return

    # Exportar a Excel
    output_dir = r"c:\Users\YARBED\Desktop\Entornos\reportes_almayor"
    os.makedirs(output_dir, exist_ok=True)

    tipo_tag = tipo_cliente.replace(' ', '_')
    excel_path = os.path.join(output_dir, f"Clientes_{tipo_tag}_Periodo_{FECHA_INICIO_PERIODO}.xlsx")

    print(f"\nGenerando Excel en: {excel_path} ...")
    exportar_excel(df, excel_path, tipo_cliente, vendedores_sel, provincias_sel, FECHA_INICIO_PERIODO)

    duracion = time.time() - inicio_tiempo
    print(f"\n[OK] Reporte generado en {duracion:.2f} segundos.")
    print(f"     Ruta: {excel_path}")


if __name__ == "__main__":
    generar_reporte()
